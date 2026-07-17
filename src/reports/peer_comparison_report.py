"""
src/reports/peer_comparison_report.py
Generates output/peer_comparison.xlsx — one sheet per peer group (11 sheets),
colour-coded percentile ranks, benchmark company highlighting, and median rows.
"""

import sqlite3
import logging
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl.styles import PatternFill, Font, Alignment

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).parent.parent.parent
DB_PATH = BASE_DIR / "data" / "database.sqlite"
OUTPUT_PATH = BASE_DIR / "output" / "peer_comparison.xlsx"

# Metrics to include (same 10 as peer.py)
METRICS = [
    {"label": "ROE (%)", "column": "return_on_equity_pct", "metric_key": "ROE"},
    {"label": "ROCE (%)", "column": "roce_percentage", "metric_key": "ROCE"},
    {"label": "NPM (%)", "column": "net_profit_margin_pct", "metric_key": "Net Profit Margin"},
    {"label": "D/E", "column": "debt_to_equity", "metric_key": "D/E"},
    {"label": "FCF (Cr)", "column": "free_cash_flow_cr", "metric_key": "FCF"},
    {"label": "PAT CAGR 5yr (%)", "column": "pat_cagr_5yr", "metric_key": "PAT CAGR 5yr"},
    {"label": "Rev CAGR 5yr (%)", "column": "revenue_cagr_5yr", "metric_key": "Revenue CAGR 5yr"},
    {"label": "EPS CAGR 5yr (%)", "column": "eps_cagr_5yr", "metric_key": "EPS CAGR 5yr"},
    {"label": "Interest Coverage", "column": "interest_coverage", "metric_key": "Interest Coverage"},
    {"label": "Asset Turnover", "column": "asset_turnover", "metric_key": "Asset Turnover"},
]

# Cell styling
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GREEN_FONT = Font(color="006100")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
YELLOW_FONT = Font(color="9C6500")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(color="9C0006")
GOLD_FILL = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
GOLD_FONT = Font(bold=True)
MEDIAN_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
MEDIAN_FONT = Font(italic=True, color="1F4E79")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)


def _load_peer_data(year: int = 2024) -> pd.DataFrame:
    """Load peer group membership with financial ratios and percentile ranks."""
    metric_cols = ", ".join([f"fr.{m['column']}" for m in METRICS])

    query = f"""
        SELECT
            pg.peer_group_name,
            pg.ticker,
            c.company_name,
            fr.composite_quality_score,
            {metric_cols}
        FROM peer_groups pg
        JOIN companies c ON pg.ticker = c.ticker
        LEFT JOIN financial_ratios fr ON pg.ticker = fr.ticker AND fr.year = :year
        ORDER BY pg.peer_group_name, pg.ticker
    """
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql(query, conn, params={"year": year})

    # Load percentile ranks and pivot to wide format
    pct_query = """
        SELECT ticker, peer_group_name, metric, percentile_rank
        FROM peer_percentiles
        WHERE year = :year
    """
    pct_df = pd.read_sql(pct_query, conn, params={"year": year})
    conn.close()

    # Force numeric
    for m in METRICS:
        col = m["column"]
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if "composite_quality_score" in df.columns:
        df["composite_quality_score"] = pd.to_numeric(df["composite_quality_score"], errors="coerce")

    # Pivot percentiles: (ticker, peer_group_name) → columns like "ROE_pctile"
    if not pct_df.empty:
        pct_pivot = pct_df.pivot_table(
            index=["ticker", "peer_group_name"],
            columns="metric",
            values="percentile_rank",
        ).reset_index()

        # Rename columns to _pctile suffix
        rename_map = {}
        for col in pct_pivot.columns:
            if col not in ["ticker", "peer_group_name"]:
                rename_map[col] = f"{col}_pctile"
        pct_pivot = pct_pivot.rename(columns=rename_map)

        df = df.merge(pct_pivot, on=["ticker", "peer_group_name"], how="left")

    return df


def generate_peer_comparison_report(year: int = 2024, output_path: Path = OUTPUT_PATH) -> None:
    """
    Generate peer_comparison.xlsx with 11 sheets — one per peer group.
    Each sheet has raw metric values + percentile rank columns, colour-coded,
    with benchmark company highlighted in gold and a median summary row.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    df = _load_peer_data(year)

    if df.empty:
        logger.warning("No peer data found — aborting report generation.")
        return

    # Build column order for each sheet
    base_cols = ["ticker", "company_name", "composite_quality_score"]
    metric_cols = []
    for m in METRICS:
        metric_cols.append(m["column"])
        pctile_col = f"{m['metric_key']}_pctile"
        metric_cols.append(pctile_col)

    all_cols = base_cols + metric_cols

    with pd.ExcelWriter(str(output_path), engine="openpyxl") as writer:
        for group_name, group_df in df.groupby("peer_group_name"):
            # Sort by composite quality score descending
            group_df = group_df.sort_values(
                "composite_quality_score", ascending=False
            ).reset_index(drop=True)

            # Select columns that exist
            cols_present = [c for c in all_cols if c in group_df.columns]
            export_df = group_df[cols_present].copy()

            # Build column header labels
            header_map = {"ticker": "Ticker", "company_name": "Company", "composite_quality_score": "CQS"}
            for m in METRICS:
                header_map[m["column"]] = m["label"]
                pctile_col = f"{m['metric_key']}_pctile"
                header_map[pctile_col] = f"{m['label']} Rank"
            export_df = export_df.rename(columns=header_map)

            # Compute median row
            median_row = {}
            for col in export_df.columns:
                if col in ["Ticker", "Company"]:
                    median_row[col] = "MEDIAN" if col == "Ticker" else ""
                else:
                    numeric_vals = pd.to_numeric(export_df[col], errors="coerce")
                    median_row[col] = round(numeric_vals.median(), 4) if numeric_vals.notna().any() else None

            median_df = pd.DataFrame([median_row])
            export_df = pd.concat([export_df, median_df], ignore_index=True)

            # Identify benchmark company (highest CQS, first data row)
            benchmark_row_idx = 0  # 0-indexed in the data, row 2 in Excel (header=1)

            # Write sheet
            sheet_name = group_name[:31]  # Excel 31-char limit
            export_df.to_excel(writer, sheet_name=sheet_name, index=False)

            # Style the sheet
            ws = writer.sheets[sheet_name]

            # Header styling
            for col_idx in range(1, len(export_df.columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", wrap_text=True)

            # Find percentile rank column indices (1-indexed)
            pctile_col_indices = []
            for col_idx, col_name in enumerate(export_df.columns, start=1):
                if col_name.endswith(" Rank"):
                    pctile_col_indices.append(col_idx)

            num_data_rows = len(export_df) - 1  # Exclude median row
            median_excel_row = num_data_rows + 2  # +1 for header, +1 for 1-indexing

            # Apply colour-coding to percentile rank cells
            for col_idx in pctile_col_indices:
                for row_idx in range(2, median_excel_row):  # Data rows only
                    cell = ws.cell(row=row_idx, column=col_idx)
                    try:
                        val = float(cell.value) if cell.value is not None else None
                    except (ValueError, TypeError):
                        val = None

                    if val is not None:
                        if val >= 0.75:
                            cell.fill = GREEN_FILL
                            cell.font = GREEN_FONT
                        elif val <= 0.25:
                            cell.fill = RED_FILL
                            cell.font = RED_FONT
                        else:
                            cell.fill = YELLOW_FILL
                            cell.font = YELLOW_FONT

            # Highlight benchmark company row (gold background)
            benchmark_excel_row = benchmark_row_idx + 2  # +1 header, +1 for 1-index
            for col_idx in range(1, len(export_df.columns) + 1):
                cell = ws.cell(row=benchmark_excel_row, column=col_idx)
                # Only apply gold to non-percentile columns (don't override colour coding)
                if col_idx not in pctile_col_indices:
                    cell.fill = GOLD_FILL
                    cell.font = GOLD_FONT

            # Style median row
            for col_idx in range(1, len(export_df.columns) + 1):
                cell = ws.cell(row=median_excel_row, column=col_idx)
                cell.fill = MEDIAN_FILL
                cell.font = MEDIAN_FONT

            # Auto-fit column widths (approximate)
            for col_idx in range(1, len(export_df.columns) + 1):
                max_len = max(
                    len(str(ws.cell(row=r, column=col_idx).value or ""))
                    for r in range(1, ws.max_row + 1)
                )
                ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 3, 22)

            logger.info(f"  Sheet '{sheet_name}': {num_data_rows} companies + median row")

    logger.info(f"peer_comparison.xlsx saved to {output_path}")


if __name__ == "__main__":
    generate_peer_comparison_report(2024)
